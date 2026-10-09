"""The files an official list is read from: what a list module declares about each
(`ListFile`), fetching it once into a local cache and checking its hash, and turning it into
lines of text with one renderer per format. A CSV or spreadsheet is one line per row with the
header first (a workbook's named sheets one after another, each with its own header), JSON is
one line per record, a web page is its visible text; a ZIP names the member to use. The text is
what the snapshot stores, so a citation's line can be checked against it.


A file is `fetched` (the URL is the file: the loader downloads it and checks its pinned hash)
or `manual` (a site that blocks scripts or only offers an interactive export: a person obtains
the file by the module's instructions and drops it in the cache, and no hash is pinned since
the next export differs). It is a file a list module reads, not the graph's `Source`, which is
a web page carrying a procurement signal.
"""

import csv
import io
import json
import logging
import re
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from enum import StrEnum
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath
from typing import IO, Any
from urllib.parse import unquote, urlsplit

import httpx
import openpyxl

from public_atlas.integrations.parse import Parser
from public_atlas.modules.evidence.service import content_hash
from public_atlas.modules.imports.models import Retrieval

logger = logging.getLogger(__name__)

# Between the cells of a table's line.
CELL_SEPARATOR = " | "
DOWNLOAD_TIMEOUT = 120.0
# Tags whose text a reader does not see.
_HIDDEN = frozenset({"script", "style", "noscript", "template", "head"})
# Tags that start a new line of visible text.
_BLOCKS = frozenset(
    {
        "p",
        "div",
        "br",
        "li",
        "tr",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "section",
        "article",
        "header",
        "footer",
        "table",
        "ul",
        "ol",
        "dt",
        "dd",
        "blockquote",
        "pre",
        "hr",
    }
)


class Format(StrEnum):
    CSV = "csv"
    SPREADSHEET = "spreadsheet"
    JSON = "json"
    HTML = "html"
    # Parsed by the parser the loader is given (Docling where it is installed).
    PDF = "pdf"


MEDIA_TYPES = {
    Format.CSV: "text/csv",
    Format.SPREADSHEET: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    Format.JSON: "application/json",
    Format.HTML: "text/html",
    Format.PDF: "application/pdf",
}
ZIP_MEDIA_TYPE = "application/zip"
TABLE_FORMATS = frozenset({Format.CSV, Format.SPREADSHEET})


SHA256 = re.compile(r"^[0-9a-f]{64}$")


class ListFileError(Exception):
    """A list's file cannot be read as the module declares it: it is missing, its hash is not
    the recorded one, it is too short, or its format is not one the loader renders."""


@dataclass(frozen=True, slots=True)
class ListFile:
    """One file a list module reads: where it comes from, how it is obtained and how to read it.
    For a fetched file the hash pins the release; a new release means a new hash and a review of
    the diff it makes. A manual file pins no hash: `instructions` say how a person gets it and
    `min_rows` catches a short export."""

    # Unique within the list; `official_lists.name` is `<list>/<name>`.
    name: str
    title: str
    # The file itself when fetched; the page the instructions start from when manual.
    url: str
    format: Format
    sha256: str | None = None
    retrieval: Retrieval = Retrieval.FETCHED
    # Manual files: the exact steps (the page, the tab, the object, the menu item, the expected
    # row count), as `lists manifest` prints them and as the error names them when the file is
    # not in the cache.
    instructions: str = ""
    # Manual tables: fewer rows than this is a wrong or truncated export.
    min_rows: int = 0
    # The file inside the ZIP at `url`, when it is one.
    member: str | None = None
    # Spreadsheets: the worksheets to read, by name and in order, each a table of its own in
    # the text (its header, then its rows; `Row.sheet` says which); None reads the first. With
    # several, the kept columns are those each sheet has, and every one must be in some sheet.
    # Two lists that read two sheets of one workbook share one `ListFile`: a file's bytes have
    # one stored text, so the sheets are rendered together.
    sheets: tuple[str, ...] | None = None

    # Text formats only.
    encoding: str = "utf-8-sig"
    # CSV: what separates the cells ("|" for the Census Bureau's code files).
    delimiter: str = ","
    # Table formats: the row the header is on, counted from 1. The rows before it (a title, a
    # note) are kept in the text and are no rows, so a line number stays the file's row number.
    header_row: int = 1
    # Table formats: keep only these columns, in this order. For a wide file read for a few of
    # its columns; the stored text is the kept columns.
    columns: tuple[str, ...] | None = None
    # Table formats: one line per distinct row of the kept columns, first seen first. For a file
    # that repeats the rows of interest once per finer unit.
    distinct: bool = False
    # Manual files: the name the file has in the cache when its URL does not end in one.
    filename_override: str | None = None
    # The file sits on a host that is not its publisher's own (a code host such as
    # raw.githubusercontent.com, where anyone publishes). The loader admits the host as a
    # platform, fetchable and never trusted, instead of trusting it as the list's domain.
    shared_host: bool = False

    def __post_init__(self) -> None:
        if self.header_row < 1:
            raise ValueError(f"{self.name}: the header row is counted from 1")
        if self.sheets is not None and (self.format is not Format.SPREADSHEET or not self.sheets):
            raise ValueError(f"{self.name}: `sheets` names a spreadsheet's worksheets")

        if self.retrieval is Retrieval.FETCHED:
            if self.sha256 is None or not SHA256.match(self.sha256):
                raise ValueError(f"{self.name}: a fetched file needs its sha256 pinned")
            if self.instructions:
                raise ValueError(f"{self.name}: a fetched file has no instructions")
        else:
            if self.sha256 is not None:
                raise ValueError(f"{self.name}: a manual file pins no hash")
            if not self.instructions.strip():
                raise ValueError(f"{self.name}: a manual file needs its instructions")
            if not self.filename:
                raise ValueError(f"{self.name}: a manual file needs a file name")

    @property
    def is_manual(self) -> bool:
        return self.retrieval is Retrieval.MANUAL

    @property
    def filename(self) -> str:
        """The file's own name: from its URL, or as the module names it."""
        if self.filename_override is not None:
            return self.filename_override
        name = PurePosixPath(unquote(urlsplit(self.url).path)).name
        return name or self.name

    @property
    def media_type(self) -> str:
        return ZIP_MEDIA_TYPE if self.member else MEDIA_TYPES[self.format]

    def cache_path(self, cache_dir: Path) -> Path:
        """Where the file sits in the cache: `<name>-<hash prefix><suffix>` for a fetched file,
        so a new release sits beside the old one, and `<name><suffix>` for a manual one."""
        suffix = PurePosixPath(self.filename).suffix
        if self.sha256 is None:
            return cache_dir / f"{self.name}{suffix}"
        return cache_dir / f"{self.name}-{self.sha256[:12]}{suffix}"


@dataclass(frozen=True, slots=True)
class Row:
    """One row of a table, with the line of the text it is and, in a workbook, its sheet."""

    line: int
    cells: dict[str, str]
    sheet: str = ""

    def __getitem__(self, column: str) -> str:
        return self.cells[column].strip()

    def get(self, column: str, default: str = "") -> str:
        return self.cells.get(column, default).strip()


@dataclass(frozen=True, slots=True)
class OpenedFile:
    """A file as `entries()` receives it: its bytes as fetched, its text as lines, and the
    rows (a table) or the parsed object (JSON) the lines were made from."""

    file: ListFile
    data: bytes
    sha256: str
    lines: list[str]
    rows: list[Row] = field(default_factory=list)
    document: Any = None

    @property
    def text(self) -> str:
        return "\n".join(self.lines)

    def line(self, number: int) -> str:
        """The text's line `number`, counted from 1."""
        return self.lines[number - 1]


# --- Fetching ---


def fetch(file: ListFile, cache_dir: Path) -> bytes:
    """The file's bytes. A fetched file comes from the cache when it is there with the right
    hash, else it is downloaded, checked and cached. A manual file is only ever read from the
    cache: when it is not there, the error says how to get it and where to put it."""
    path = file.cache_path(cache_dir)
    if file.is_manual:
        if not path.exists():
            raise ListFileError(
                f"{file.name}: the file is obtained by hand and is not in the cache. "
                f"Put it at {path}. {file.instructions.strip()}"
            )
        return path.read_bytes()
    if path.exists():
        data = path.read_bytes()
        if content_hash(data) == file.sha256:
            return data
        logger.warning("%s does not match %s's hash; fetching again", path, file.name)
    logger.info("fetching %s", file.url)
    try:
        response = httpx.get(file.url, follow_redirects=True, timeout=DOWNLOAD_TIMEOUT)
        response.raise_for_status()
    except httpx.HTTPError as exc:
        raise ListFileError(f"{file.name}: cannot fetch {file.url}: {exc}") from None
    data = response.content
    found = content_hash(data)
    if found != file.sha256:
        raise ListFileError(
            f"{file.name}: {file.url} has sha256 {found}, not the recorded {file.sha256}. "
            "A new release needs its hash recorded in the list module and the diff reviewed"
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return data


def open_file(file: ListFile, cache_dir: Path, *, parser: Parser | None = None) -> OpenedFile:
    return render(file, fetch(file, cache_dir), parser=parser)


# --- Rendering ---


def render(file: ListFile, data: bytes, *, parser: Parser | None = None) -> OpenedFile:
    """The file as lines of text, with its rows or parsed object. A PDF goes through `parser`,
    one line of text per line of its pages. A manual file's hash is whatever it is."""
    digest = content_hash(data)
    if file.sha256 is not None and digest != file.sha256:
        raise ListFileError(f"{file.name}: sha256 {digest} is not the recorded {file.sha256}")
    with _member(file, data) as stream:
        if file.format in TABLE_FORMATS:
            tables: list[tuple[str, Iterable[list[str]]]]
            if file.format is Format.CSV:
                tables = [("", _csv_rows(stream, file.encoding, file.delimiter))]
            else:
                tables = _sheet_rows(file, stream)
            lines, rows = _tables(file, tables)
            return OpenedFile(file=file, data=data, sha256=digest, lines=lines, rows=rows)

        if file.format is Format.JSON:
            document = json.loads(stream.read().decode(file.encoding))
            return OpenedFile(
                file=file, data=data, sha256=digest, lines=_json_lines(document), document=document
            )
        if file.format is Format.HTML:
            text = stream.read().decode(file.encoding, errors="replace")
            return OpenedFile(file=file, data=data, sha256=digest, lines=visible_lines(text))
        if file.format is Format.PDF:
            if parser is None:
                raise ListFileError(f"{file.name}: a PDF list needs a parser")
            document = parser.parse(stream.read(), file.filename)
            lines = [line for page in document.pages for line in page.splitlines()]
            return OpenedFile(file=file, data=data, sha256=digest, lines=lines)
    raise ListFileError(f"{file.name}: no renderer for {file.format}")  # pragma: no cover


def _member(file: ListFile, data: bytes) -> IO[bytes]:
    """The bytes to read: the ZIP member, streamed, or the file itself."""
    if file.member is None:
        return io.BytesIO(data)
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        return archive.open(file.member)
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ListFileError(f"{file.name}: no member {file.member!r}: {exc}") from None


def _csv_rows(stream: IO[bytes], encoding: str, delimiter: str) -> Iterator[list[str]]:
    reader = csv.reader(
        io.TextIOWrapper(stream, encoding=encoding, newline=""), delimiter=delimiter
    )
    yield from reader


def _sheet_rows(file: ListFile, stream: IO[bytes]) -> list[tuple[str, Iterable[list[str]]]]:
    """The worksheets the file names, or the first, each by its title with every cell as
    text."""
    workbook = openpyxl.load_workbook(io.BytesIO(stream.read()), read_only=True, data_only=True)
    try:
        if file.sheets is None:
            sheets = [workbook.worksheets[0]]
        else:
            for name in file.sheets:
                if name not in workbook.sheetnames:
                    raise ListFileError(f"{file.name}: the workbook has no sheet {name!r}")
            sheets = [workbook[name] for name in file.sheets]
        return [
            (
                # A row says which sheet it is from when the file names its sheets.
                sheet.title if file.sheets is not None else "",
                [
                    ["" if cell is None else str(cell) for cell in row]
                    for row in sheet.iter_rows(values_only=True)
                ],
            )
            for sheet in sheets
        ]
    finally:
        workbook.close()


def _tables(
    file: ListFile, tables: Iterable[tuple[str, Iterable[list[str]]]]
) -> tuple[list[str], list[Row]]:
    """Lines and rows from a file's tables: a CSV's one, a workbook's sheets one after another.
    With several sheets, each keeps the columns it has, and a column no sheet has is an error.
    Fewer rows in all than the file's `min_rows` is a wrong or truncated export."""
    lines: list[str] = []
    rows: list[Row] = []
    several = file.sheets is not None and len(file.sheets) > 1
    found: set[str] = set()
    for sheet, raw in tables:
        found |= _table(file, raw, sheet, lines, rows, strict=not several)
    if file.columns is not None:
        missing = [column for column in file.columns if column not in found]
        if missing:
            raise ListFileError(f"{file.name}: no sheet has the columns {missing}")
    if len(rows) < file.min_rows:
        raise ListFileError(
            f"{file.name}: the table has {len(rows)} rows, fewer than the {file.min_rows} "
            "expected: a wrong or truncated export"
        )
    return lines, rows


def _table(  # noqa: PLR0913 - the accumulators and what fills them
    file: ListFile,
    raw: Iterable[list[str]],
    sheet: str,
    lines: list[str],
    rows: list[Row],
    *,
    strict: bool,
) -> set[str]:
    """One table onto `lines` and `rows`: the rows before the header as they are, the header,
    then each row's cells joined, the columns kept as the file says (every one when `strict`,
    else those the table has) and repeats dropped when it asks. The columns kept."""
    rows_iter = iter(raw)
    for _ in range(file.header_row - 1):
        before = next(rows_iter, None)
        if before is None:
            break
        lines.append(CELL_SEPARATOR.join(cell.strip() for cell in before))
    header = next(rows_iter, None)
    if header is None:
        raise ListFileError(f"{file.name}: the table has no row {file.header_row} to head it")
    header = [cell.strip() for cell in header]
    if file.columns is not None:
        wanted = [column for column in file.columns if strict or column in header]
        try:
            keep = [header.index(column) for column in wanted]
        except ValueError as exc:
            raise ListFileError(f"{file.name}: the table has no column {exc}") from None
        header = wanted
    else:
        keep = None
    lines.append(CELL_SEPARATOR.join(header))
    seen: set[tuple[str, ...]] = set()
    for raw_cells in rows_iter:
        cells = raw_cells if keep is None else [_cell(raw_cells, index) for index in keep]
        if file.distinct:
            key = tuple(cells)
            if key in seen:
                continue
            seen.add(key)
        lines.append(CELL_SEPARATOR.join(cells))
        rows.append(Row(line=len(lines), cells=dict(zip(header, cells, strict=False)), sheet=sheet))
    return set(header)


def _cell(cells: list[str], index: int) -> str:
    return cells[index] if index < len(cells) else ""


def _json_lines(document: Any) -> list[str]:  # noqa: ANN401 - JSON is any value
    records = document if isinstance(document, list) else [document]
    return [json.dumps(record, ensure_ascii=False, sort_keys=True) for record in records]


class _VisibleText(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.lines: list[str] = []
        self._current: list[str] = []
        self._hidden = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:  # noqa: ARG002
        if tag in _HIDDEN:
            self._hidden += 1
        if tag in _BLOCKS:
            self._flush()

    def handle_endtag(self, tag: str) -> None:
        if tag in _HIDDEN and self._hidden:
            self._hidden -= 1
        if tag in _BLOCKS:
            self._flush()

    def handle_data(self, data: str) -> None:
        if not self._hidden:
            self._current.append(data)

    def _flush(self) -> None:
        text = " ".join("".join(self._current).split())
        if text:
            self.lines.append(text)
        self._current = []

    def close(self) -> None:
        super().close()
        self._flush()


def visible_lines(html: str) -> list[str]:
    """A page's visible text, one line per block."""
    parser = _VisibleText()
    parser.feed(html)
    parser.close()
    return parser.lines
